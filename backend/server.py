from dotenv import load_dotenv
load_dotenv()

import os
import hmac
import logging
from datetime import datetime, timezone, timedelta
from typing import Optional, List

import bcrypt
import jwt
from bson import ObjectId
from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, BackgroundTasks
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

JWT_ALGORITHM = "HS256"
JWT_SECRET = os.environ["JWT_SECRET"]
IST = timezone(timedelta(hours=5, minutes=30))

app = FastAPI()
api_router = APIRouter(prefix="/api")
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("mukman")


def now():
    return datetime.now(timezone.utc)


def now_iso():
    return now().isoformat()


def hash_password(p: str) -> str:
    return bcrypt.hashpw(p.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(p: str, h: str) -> bool:
    return bcrypt.checkpw(p.encode("utf-8"), h.encode("utf-8"))


def create_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": now() + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


def serialize(doc):
    if not doc:
        return None
    doc = dict(doc)
    doc["id"] = str(doc.pop("_id"))
    doc.pop("password_hash", None)
    for k in ("reports_to", "approver_id", "user_id", "current_approver_id", "current_owner_id", "reviewed_by", "updated_by", "also_reports_to"):
        if isinstance(doc.get(k), ObjectId):
            doc[k] = str(doc[k])
        elif isinstance(doc.get(k), list):
            doc[k] = [str(x) if isinstance(x, ObjectId) else x for x in doc[k]]
    for list_key in ("chain", "history"):
        if isinstance(doc.get(list_key), list):
            for entry in doc[list_key]:
                if isinstance(entry, dict):
                    for kk, vv in list(entry.items()):
                        if isinstance(vv, ObjectId):
                            entry[kk] = str(vv)
    return doc


async def get_current_user(request: Request):
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")
    user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
    if not user or not user.get("active", True):
        raise HTTPException(401, "User not found")
    return serialize(user)


async def require_admin(user=Depends(get_current_user)):
    if user["role"] != "admin":
        raise HTTPException(403, "Admin access only")
    return user


async def notify(user_id, ntype, title, message, channel="in_app", popup=False):
    await db.notifications.insert_one({
        "user_id": ObjectId(str(user_id)), "type": ntype, "title": title, "message": message,
        "channel": channel, "popup": popup, "read": False, "created_at": now_iso(),
    })


# ---------------- Auth ----------------

class LoginReq(BaseModel):
    identifier: str
    password: str


@api_router.post("/auth/login")
async def login(req: LoginReq, response: Response):
    ident = req.identifier.strip()
    user = await db.users.find_one({"$or": [{"email": ident.lower()}, {"phone": ident}]})
    if not user or not verify_password(req.password, user["password_hash"]):
        raise HTTPException(401, "Invalid credentials")
    if not user.get("active", True):
        raise HTTPException(403, "Account deactivated. Contact admin.")
    token = create_token(str(user["_id"]))
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=7 * 86400, path="/")
    return serialize(user)


@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api_router.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return user


class ChangePwReq(BaseModel):
    current_password: str
    new_password: str


@api_router.post("/auth/change-password")
async def change_password(req: ChangePwReq, user=Depends(get_current_user)):
    if len(req.new_password) < 6:
        raise HTTPException(400, "New password must be at least 6 characters")
    doc = await db.users.find_one({"_id": ObjectId(user["id"])})
    if not doc or not verify_password(req.current_password, doc["password_hash"]):
        raise HTTPException(400, "Current password is incorrect")
    await db.users.update_one({"_id": doc["_id"]}, {"$set": {"password_hash": hash_password(req.new_password), "must_change_password": False}})
    return {"ok": True}


# ---------------- Content CMS ----------------

@api_router.get("/content")
async def get_content(user=Depends(get_current_user)):
    docs = await db.content.find({}).to_list(100)
    return {d["key"]: serialize(d) for d in docs}


class ContentUpdate(BaseModel):
    title: str
    body: str
    stats: Optional[List[dict]] = None


@api_router.put("/admin/content/{key}")
async def update_content(key: str, req: ContentUpdate, admin=Depends(require_admin)):
    update = {"title": req.title, "body": req.body, "updated_at": now_iso(), "updated_by": ObjectId(admin["id"])}
    if req.stats is not None:
        update["stats"] = req.stats
    await db.content.update_one({"key": key}, {"$set": update}, upsert=True)
    doc = await db.content.find_one({"key": key})
    return serialize(doc)


class EnhanceReq(BaseModel):
    text: str
    section: str = "general"
    tone: str = "corporate"


@api_router.post("/admin/ai-enhance")
async def ai_enhance(req: EnhanceReq, admin=Depends(require_admin)):
    from emergentintegrations.llm.chat import LlmChat, UserMessage

    system = (
        "You are the corporate communications writer for Mukman Technologies' internal employee portal. "
        "Mukman Group is a mobility holding company (SchoolCab: school commute, Cabzee: fixed-route office commute, "
        "Mukman Advisory: financial services). Rewrite the rough draft into polished, professional, inspiring intranet content. "
        "Expand thin one-liners into rich content: add context, explain terms, and why it matters to employees. "
        "Keep every factual number, name and figure EXACTLY as given. Never invent new numbers. "
        "If the section is 'valuation', additionally sense-check that the market-size framing is coherent and clearly explain "
        "the valuation figures chosen. Use short paragraphs. Return ONLY the enhanced text, no preamble."
    )
    chat = LlmChat(
        api_key=os.environ["EMERGENT_LLM_KEY"],
        session_id=f"enhance-{admin['id']}-{int(now().timestamp())}",
        system_message=system,
    ).with_model("openai", "gpt-5.4")
    msg = UserMessage(text=f"Section: {req.section}\nTone: {req.tone}\n\nDraft text:\n{req.text}")
    enhanced = await chat.send_message(msg)
    return {"enhanced": str(enhanced)}


# ---------------- Directory / Org chart ----------------

@api_router.get("/org-chart")
async def org_chart(user=Depends(get_current_user)):
    users = await db.users.find({"active": True}).to_list(500)
    people = [
        {"id": str(u["_id"]), "name": u["name"], "title": u.get("title", ""), "department": u.get("department", ""),
         "role": u["role"], "reports_to": str(u["reports_to"]) if u.get("reports_to") else None,
         "also_reports_to": [str(x) for x in u.get("also_reports_to", [])],
         "esop_applicable": u.get("esop_applicable", False)}
        for u in users
    ]
    founder = next((u for u in users if u["role"] == "admin" and not u.get("reports_to")), None)
    virtual = []
    async for v in db.org_nodes.find({}):
        pid = None
        if v.get("parent_email") == "founder":
            pid = str(founder["_id"]) if founder else None
        elif v.get("parent_email"):
            m = await db.users.find_one({"email": v["parent_email"]})
            pid = str(m["_id"]) if m else None
        virtual.append({"id": v["key"], "name": v["name"], "title": v.get("title", ""),
                        "kind": v.get("kind", "open"), "reports_to": pid, "virtual": True})
    return {"people": people, "virtual": virtual}


# ---------------- Timesheets ----------------

class TaskEntry(BaseModel):
    category: str
    hours: float = 0
    note: str = ""


class TimesheetReq(BaseModel):
    date: str
    status: str
    tasks: List[TaskEntry] = []
    idle_hours: float = 0
    idle_reason: str = ""
    achievements: str = ""


@api_router.post("/timesheets")
async def submit_timesheet(req: TimesheetReq, user=Depends(get_current_user)):
    today = now().astimezone(IST).date().isoformat()
    if req.date > today:
        raise HTTPException(400, "Cannot fill timesheet for a future date")
    if req.status not in ("present", "absent", "leave"):
        raise HTTPException(400, "Invalid status")
    tasks = [t.model_dump() for t in req.tasks if t.hours > 0 or t.note.strip()]
    total_hours = round(sum(t["hours"] for t in tasks), 2)
    if req.status == "present" and total_hours + req.idle_hours > 24:
        raise HTTPException(400, "Total hours (incl. idle) cannot exceed 24 in a day")
    if req.idle_hours > 0 and not req.idle_reason.strip():
        raise HTTPException(400, "Please add a reason for idle hours")
    existing = await db.timesheets.find_one({"user_id": ObjectId(user["id"]), "date": req.date})
    if existing and existing.get("approval_status") == "approved":
        raise HTTPException(400, "This day is already approved and locked")
    approver = user.get("reports_to")
    doc = {
        "user_id": ObjectId(user["id"]), "user_name": user["name"], "date": req.date,
        "status": req.status, "tasks": tasks, "idle_hours": req.idle_hours, "idle_reason": req.idle_reason,
        "hours": total_hours, "total_hours": total_hours, "achievements": req.achievements,
        "approval_status": "pending" if approver else "approved",
        "approver_id": ObjectId(approver) if approver else None,
        "comments": "", "created_at": now_iso(), "reviewed_at": None,
    }
    if existing:
        await db.timesheets.update_one({"_id": existing["_id"]}, {"$set": doc})
        tid = str(existing["_id"])
    else:
        res = await db.timesheets.insert_one(doc)
        tid = str(res.inserted_id)
    if approver:
        await notify(approver, "approval_request", f"Timesheet approval: {user['name']}",
                     f"{user['name']} submitted a timesheet for {req.date} ({req.status}). Please review.")
    return {"id": tid, **{k: (str(v) if isinstance(v, ObjectId) else v) for k, v in doc.items()}}


@api_router.get("/timesheets/my")
async def my_timesheets(month: Optional[str] = None, user=Depends(get_current_user)):
    q = {"user_id": ObjectId(user["id"])}
    if month:
        q["date"] = {"$regex": f"^{month}"}
    docs = await db.timesheets.find(q).sort("date", -1).to_list(400)
    return [serialize(d) for d in docs]


@api_router.get("/timesheets/approvals")
async def timesheet_approvals(user=Depends(get_current_user)):
    q = {"approver_id": ObjectId(user["id"]), "approval_status": "pending"}
    docs = await db.timesheets.find(q).sort("date", -1).to_list(400)
    return [serialize(d) for d in docs]


async def _review_timesheet(tid, action, comment, user):
    ts = await db.timesheets.find_one({"_id": ObjectId(tid)})
    if not ts:
        raise HTTPException(404, "Timesheet not found")
    if str(ts.get("approver_id")) != user["id"] and user["role"] != "admin":
        raise HTTPException(403, "Not the approver for this entry")
    if ts.get("approval_status") != "pending":
        raise HTTPException(400, "Already reviewed")
    await db.timesheets.update_one({"_id": ts["_id"]}, {"$set": {
        "approval_status": action, "comments": comment, "reviewed_by": ObjectId(user["id"]), "reviewed_at": now_iso()}})
    verb = "approved" if action == "approved" else "rejected"
    await notify(ts["user_id"], "status_update", f"Timesheet {verb}: {ts['date']}",
                 f"Your timesheet for {ts['date']} was {verb} by {user['name']}." + (f" Comment: {comment}" if comment else ""))
    return {"ok": True}


class CommentReq(BaseModel):
    comment: str = ""


@api_router.post("/timesheets/{tid}/approve")
async def approve_timesheet(tid: str, req: CommentReq, user=Depends(get_current_user)):
    return await _review_timesheet(tid, "approved", req.comment, user)


@api_router.post("/timesheets/{tid}/reject")
async def reject_timesheet(tid: str, req: CommentReq, user=Depends(get_current_user)):
    return await _review_timesheet(tid, "rejected", req.comment, user)


@api_router.get("/timesheets/export")
async def export_timesheets(period: str = "month", start: Optional[str] = None, end: Optional[str] = None, user=Depends(get_current_user)):
    today = now().astimezone(IST).date()
    if period == "week":
        s = (today - timedelta(days=today.weekday())).isoformat()
        e = today.isoformat()
    elif period == "month":
        s = today.replace(day=1).isoformat()
        e = today.isoformat()
    elif period == "ytd":
        s = today.replace(month=1, day=1).isoformat()
        e = today.isoformat()
    elif period == "custom" and start and end and start <= end:
        s, e = start, end
    else:
        raise HTTPException(400, "Invalid period or date range")

    users = await db.users.find({"active": True}).to_list(500)
    children = {}
    for u in users:
        if u.get("reports_to"):
            children.setdefault(str(u["reports_to"]), []).append(u)
    scope = []
    stack = list(children.get(user["id"], []))
    while stack:
        u = stack.pop()
        scope.append(u)
        stack.extend(children.get(str(u["_id"]), []))
    if user["role"] == "admin":
        have = {u["_id"] for u in scope}
        for u in users:
            if u["_id"] not in have and str(u["_id"]) != user["id"]:
                scope.append(u)
    if not scope:
        raise HTTPException(403, "You have no reportees to export")

    ids = [u["_id"] for u in scope]
    entries = await db.timesheets.find({"user_id": {"$in": ids}, "date": {"$gte": s, "$lte": e}}).sort("date", 1).to_list(10000)

    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timesheets"
    headers = ["Employee", "Date", "Status", "Core Activity (h)", "Ad-hoc Activity (h)", "Other Activities (h)",
               "Idle (h)", "Idle Reason", "Total Utilised (h)", "Achievements / Notes", "Approval Status", "Manager Comment"]
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")
    ws.append(headers)
    for c in ws[1]:
        c.font = head_font
        c.fill = head_fill

    def cat_hours(tasks, prefix):
        return round(sum(t.get("hours", 0) for t in (tasks or []) if str(t.get("category", "")).startswith(prefix)), 2)

    for en in entries:
        ws.append([
            en.get("user_name", ""), en.get("date", ""), str(en.get("status", "")).title(),
            cat_hours(en.get("tasks"), "Core"), cat_hours(en.get("tasks"), "Ad-hoc"), cat_hours(en.get("tasks"), "Other"),
            en.get("idle_hours", 0), en.get("idle_reason", ""), en.get("total_hours", en.get("hours", 0)),
            en.get("achievements", ""), str(en.get("approval_status", "")).title(), en.get("comments", ""),
        ])
    for i, w in enumerate([18, 12, 10, 16, 18, 18, 9, 28, 16, 36, 14, 24], 1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws2 = wb.create_sheet("Summary")
    ws2.append(["Employee", "Present Days", "Absent", "Leave", "Total Utilised (h)", "Idle (h)"])
    for c in ws2[1]:
        c.font = head_font
        c.fill = head_fill
    for u in scope:
        mine = [en for en in entries if en.get("user_id") == u["_id"]]
        ws2.append([
            u["name"],
            sum(1 for en in mine if en.get("status") == "present"),
            sum(1 for en in mine if en.get("status") == "absent"),
            sum(1 for en in mine if en.get("status") == "leave"),
            round(sum(en.get("total_hours", en.get("hours", 0)) for en in mine), 2),
            round(sum(en.get("idle_hours", 0) for en in mine), 2),
        ])
    for i, w in enumerate([18, 13, 9, 9, 18, 9], 1):
        ws2.column_dimensions[chr(64 + i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"mukman-timesheets-{period}-{s}-to-{e}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ---------------- Periodic reviews ----------------

class ReviewReq(BaseModel):
    period_type: str
    period_label: str
    content: str


@api_router.post("/reviews")
async def submit_review(req: ReviewReq, user=Depends(get_current_user)):
    if req.period_type not in ("weekly", "monthly", "quarterly", "annual"):
        raise HTTPException(400, "Invalid period type")
    approver = user.get("reports_to")
    doc = {
        "user_id": ObjectId(user["id"]), "user_name": user["name"],
        "period_type": req.period_type, "period_label": req.period_label, "content": req.content,
        "status": "pending" if approver else "approved",
        "current_approver_id": ObjectId(approver) if approver else None,
        "chain": [], "created_at": now_iso(),
    }
    res = await db.reviews.insert_one(doc)
    if approver:
        await notify(approver, "approval_request", f"{req.period_type.title()} review: {user['name']}",
                     f"{user['name']} submitted their {req.period_type} review ({req.period_label}). Please review.")
    return {"id": str(res.inserted_id)}


@api_router.get("/reviews/my")
async def my_reviews(user=Depends(get_current_user)):
    docs = await db.reviews.find({"user_id": ObjectId(user["id"])}).sort("created_at", -1).to_list(200)
    return [serialize(d) for d in docs]


@api_router.get("/reviews/pending")
async def pending_reviews(user=Depends(get_current_user)):
    docs = await db.reviews.find({"current_approver_id": ObjectId(user["id"]), "status": "pending"}).sort("created_at", -1).to_list(200)
    return [serialize(d) for d in docs]


@api_router.post("/reviews/{rid}/approve")
async def approve_review(rid: str, req: CommentReq, user=Depends(get_current_user)):
    r = await db.reviews.find_one({"_id": ObjectId(rid)})
    if not r:
        raise HTTPException(404, "Review not found")
    if str(r.get("current_approver_id")) != user["id"] and user["role"] != "admin":
        raise HTTPException(403, "Not the current approver")
    if r["status"] != "pending":
        raise HTTPException(400, "Already finalized")
    entry = {"approver_id": user["id"], "name": user["name"], "action": "approved", "comment": req.comment, "at": now_iso()}
    me = await db.users.find_one({"_id": ObjectId(user["id"])})
    nxt = me.get("reports_to") if me else None
    if nxt:
        await db.reviews.update_one({"_id": r["_id"]}, {"$push": {"chain": entry}, "$set": {"current_approver_id": ObjectId(nxt)}})
        await notify(nxt, "approval_request", f"{r['period_type'].title()} review escalation: {r['user_name']}",
                     f"{user['name']} approved {r['user_name']}'s {r['period_type']} review ({r['period_label']}). Final approval pending with you.")
        await notify(r["user_id"], "status_update", "Review moved up",
                     f"Your {r['period_type']} review ({r['period_label']}) was approved by {user['name']} and escalated for final approval.")
    else:
        await db.reviews.update_one({"_id": r["_id"]}, {"$push": {"chain": entry}, "$set": {"status": "approved", "current_approver_id": None}})
        await notify(r["user_id"], "status_update", "Review fully approved",
                     f"Your {r['period_type']} review ({r['period_label']}) received final approval from {user['name']}.")
    return {"ok": True}


@api_router.post("/reviews/{rid}/reject")
async def reject_review(rid: str, req: CommentReq, user=Depends(get_current_user)):
    r = await db.reviews.find_one({"_id": ObjectId(rid)})
    if not r:
        raise HTTPException(404, "Review not found")
    if str(r.get("current_approver_id")) != user["id"] and user["role"] != "admin":
        raise HTTPException(403, "Not the current approver")
    entry = {"approver_id": user["id"], "name": user["name"], "action": "rejected", "comment": req.comment, "at": now_iso()}
    await db.reviews.update_one({"_id": r["_id"]}, {"$push": {"chain": entry}, "$set": {"status": "rejected", "current_approver_id": None}})
    await notify(r["user_id"], "status_update", "Review rejected",
                 f"Your {r['period_type']} review ({r['period_label']}) was rejected by {user['name']}." + (f" Comment: {req.comment}" if req.comment else ""))
    return {"ok": True}


# ---------------- Feedback ----------------

class FeedbackReq(BaseModel):
    subject: str
    message: str


@api_router.post("/feedback")
async def submit_feedback(req: FeedbackReq, user=Depends(get_current_user)):
    owner = user.get("reports_to")
    doc = {
        "user_id": ObjectId(user["id"]), "user_name": user["name"],
        "subject": req.subject, "message": req.message, "status": "open",
        "current_owner_id": ObjectId(owner) if owner else None,
        "history": [], "created_at": now_iso(),
    }
    res = await db.feedback.insert_one(doc)
    if owner:
        await notify(owner, "approval_request", f"Feedback from {user['name']}",
                     f"New feedback: {req.subject}. Please review in the Feedback section.")
    return {"id": str(res.inserted_id)}


@api_router.get("/feedback/my")
async def my_feedback(user=Depends(get_current_user)):
    docs = await db.feedback.find({"user_id": ObjectId(user["id"])}).sort("created_at", -1).to_list(200)
    return [serialize(d) for d in docs]


@api_router.get("/feedback/inbox")
async def feedback_inbox(user=Depends(get_current_user)):
    docs = await db.feedback.find({"current_owner_id": ObjectId(user["id"]), "status": {"$ne": "resolved"}}).sort("created_at", -1).to_list(200)
    return [serialize(d) for d in docs]


@api_router.post("/feedback/{fid}/forward")
async def forward_feedback(fid: str, user=Depends(get_current_user)):
    fb = await db.feedback.find_one({"_id": ObjectId(fid)})
    if not fb:
        raise HTTPException(404, "Feedback not found")
    if str(fb.get("current_owner_id")) != user["id"] and user["role"] != "admin":
        raise HTTPException(403, "Not assigned to you")
    me = await db.users.find_one({"_id": ObjectId(user["id"])})
    nxt = me.get("reports_to") if me else None
    if not nxt:
        raise HTTPException(400, "No senior above you to forward to")
    await db.feedback.update_one({"_id": fb["_id"]}, {"$set": {"current_owner_id": ObjectId(nxt), "status": "forwarded"},
        "$push": {"history": {"by": user["name"], "action": "forwarded", "at": now_iso()}}})
    await notify(nxt, "approval_request", f"Escalated feedback: {fb['subject']}",
                 f"{user['name']} forwarded feedback from {fb['user_name']} to you.")
    return {"ok": True}


@api_router.post("/feedback/{fid}/resolve")
async def resolve_feedback(fid: str, user=Depends(get_current_user)):
    fb = await db.feedback.find_one({"_id": ObjectId(fid)})
    if not fb:
        raise HTTPException(404, "Feedback not found")
    if str(fb.get("current_owner_id")) != user["id"] and user["role"] != "admin":
        raise HTTPException(403, "Not assigned to you")
    await db.feedback.update_one({"_id": fb["_id"]}, {"$set": {"status": "resolved", "current_owner_id": None},
        "$push": {"history": {"by": user["name"], "action": "resolved", "at": now_iso()}}})
    await notify(fb["user_id"], "status_update", "Feedback resolved",
                 f"Your feedback \"{fb['subject']}\" was marked resolved by {user['name']}.")
    return {"ok": True}


# ---------------- Notifications ----------------

@api_router.get("/notifications")
async def get_notifications(user=Depends(get_current_user)):
    docs = await db.notifications.find({"user_id": ObjectId(user["id"])}).sort("created_at", -1).to_list(200)
    return [serialize(d) for d in docs]


@api_router.post("/notifications/{nid}/read")
async def read_notification(nid: str, user=Depends(get_current_user)):
    await db.notifications.update_one({"_id": ObjectId(nid), "user_id": ObjectId(user["id"])}, {"$set": {"read": True}})
    return {"ok": True}


@api_router.post("/notifications/read-all")
async def read_all_notifications(user=Depends(get_current_user)):
    await db.notifications.update_many({"user_id": ObjectId(user["id"])}, {"$set": {"read": True}})
    return {"ok": True}


# ---------------- Penalty rules ----------------

@api_router.get("/penalty-rules")
async def get_penalty_rules(user=Depends(get_current_user)):
    docs = await db.penalty_rules.find({}).to_list(50)
    return [serialize(d) for d in docs]


class PenaltyUpdate(BaseModel):
    percent: float
    description: str
    enabled: bool


@api_router.put("/admin/penalty-rules/{key}")
async def update_penalty_rule(key: str, req: PenaltyUpdate, admin=Depends(require_admin)):
    await db.penalty_rules.update_one({"key": key}, {"$set": {"percent": req.percent, "description": req.description, "enabled": req.enabled}})
    doc = await db.penalty_rules.find_one({"key": key})
    return serialize(doc)


# ---------------- Admin: employees ----------------

class EmployeeReq(BaseModel):
    name: str
    email: str
    phone: str
    password: Optional[str] = None
    title: str = ""
    department: str = ""
    reports_to: Optional[str] = None
    role: str = "employee"
    esop_applicable: bool = False
    active: bool = True


@api_router.get("/admin/employees")
async def list_employees(admin=Depends(require_admin)):
    users = await db.users.find({}).to_list(500)
    return [serialize(u) for u in users]


@api_router.post("/admin/employees")
async def create_employee(req: EmployeeReq, admin=Depends(require_admin)):
    email = req.email.strip().lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(400, "Email already exists")
    doc = {
        "name": req.name, "email": email, "phone": req.phone,
        "password_hash": hash_password(req.password or "Schoolcab@1"),
        "title": req.title, "department": req.department,
        "reports_to": ObjectId(req.reports_to) if req.reports_to else None,
        "role": req.role if req.role in ("admin", "employee") else "employee",
        "esop_applicable": req.esop_applicable, "active": req.active, "must_change_password": True, "created_at": now_iso(),
    }
    res = await db.users.insert_one(doc)
    doc["_id"] = res.inserted_id
    return serialize(doc)


@api_router.put("/admin/employees/{eid}")
async def update_employee(eid: str, req: EmployeeReq, admin=Depends(require_admin)):
    emp = await db.users.find_one({"_id": ObjectId(eid)})
    if not emp:
        raise HTTPException(404, "Employee not found")
    update = {
        "name": req.name, "email": req.email.strip().lower(), "phone": req.phone,
        "title": req.title, "department": req.department,
        "reports_to": ObjectId(req.reports_to) if req.reports_to else None,
        "role": req.role if req.role in ("admin", "employee") else "employee",
        "esop_applicable": req.esop_applicable, "active": req.active,
    }
    if req.password:
        update["password_hash"] = hash_password(req.password)
        update["must_change_password"] = True
    await db.users.update_one({"_id": emp["_id"]}, {"$set": update})
    return serialize(await db.users.find_one({"_id": emp["_id"]}))


class BroadcastReq(BaseModel):
    title: str
    message: str


@api_router.post("/admin/broadcast")
async def broadcast(req: BroadcastReq, admin=Depends(require_admin)):
    users = await db.users.find({"active": True}).to_list(500)
    for u in users:
        await notify(u["_id"], "broadcast", req.title, req.message, channel="in_app", popup=True)
    return {"sent": len(users)}


# ---------------- Cron endpoints ----------------

def verify_cron(request: Request):
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    if not token or not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(401, "Unauthorized")


async def run_timesheet_reminders(run_id: str):
    try:
        await db.cron_runs.insert_one({"run_id": run_id, "job": "timesheet", "at": now_iso()})
    except Exception:
        return
    today = now().astimezone(IST).date().isoformat()
    rules = {r["key"]: r async for r in db.penalty_rules.find({})}
    sal = rules.get("ts_delay_5d", {}).get("percent", 10)
    esop = rules.get("ts_delay_10d", {}).get("percent", 10)
    users = await db.users.find({"active": True}).to_list(500)
    for u in users:
        ts = await db.timesheets.find_one({"user_id": u["_id"], "date": today})
        if not ts:
            msg = (f"Hi {u['name']}, your timesheet for {today} is not filled yet. "
                   f"A continuous delay of 5 days will result in a {sal}% salary payout decline")
            msg += f", and a 10-day delay in a {esop}% ownership/ESOP decline." if u.get("esop_applicable") else "."
            await notify(u["_id"], "penalty_warning", f"Timesheet missing for {today}", msg, channel="whatsapp_mock")
    logger.info(f"timesheet reminders done for {today}")


async def run_weekly_review_reminders(run_id: str):
    try:
        await db.cron_runs.insert_one({"run_id": run_id, "job": "weekly_review", "at": now_iso()})
    except Exception:
        return
    d = now().astimezone(IST).date()
    label = f"{d.isocalendar()[0]}-W{d.isocalendar()[1]:02d}"
    rules = {r["key"]: r async for r in db.penalty_rules.find({})}
    pct = rules.get("weekly_review_3_missed", {}).get("percent", 10)
    users = await db.users.find({"active": True}).to_list(500)
    for u in users:
        r = await db.reviews.find_one({"user_id": u["_id"], "period_type": "weekly", "period_label": label})
        if not r:
            await notify(u["_id"], "penalty_warning", f"Weekly review due: {label}",
                         f"Hi {u['name']}, your weekly summary for {label} is due today. Missing 3 consecutive weekly summaries "
                         f"will result in a {pct}% penalty on salary or ownership as per policy.", channel="whatsapp_mock")
    logger.info(f"weekly review reminders done for {label}")


@api_router.post("/cron/timesheet-reminders")
async def cron_timesheet_reminders(request: Request, background_tasks: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    verify_cron(request)
    try:
        body = await request.json()
    except Exception:
        body = {}
    run_id = body.get("run_id") or request.headers.get("X-Webhook-Id") or f"manual-{int(now().timestamp())}"
    background_tasks.add_task(run_timesheet_reminders, run_id)
    return {"ok": True, "run_id": run_id}


@api_router.post("/cron/weekly-review-reminder")
async def cron_weekly_review(request: Request, background_tasks: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    verify_cron(request)
    try:
        body = await request.json()
    except Exception:
        body = {}
    run_id = body.get("run_id") or request.headers.get("X-Webhook-Id") or f"manual-{int(now().timestamp())}"
    background_tasks.add_task(run_weekly_review_reminders, run_id)
    return {"ok": True, "run_id": run_id}


@api_router.get("/")
async def root():
    return {"message": "Mukman Group Internal Portal API"}


# ---------------- Seed ----------------

SEED_CONTENT = {
    "vision": {
        "title": "The Mukman Vision",
        "body": "Mukman Technologies is building India's leading platform for recurring and fixed-route rides.\n\nWhile Ola and Uber solved on-demand mobility, nobody owns the rides that happen every single day — the school run, the office commute, the fixed route. We exist to take all the stress out of mobility for people who ride on a regular basis: no daily booking, no waiting for a driver, no inconsistency in who shows up.\n\nOur goal is simple — fix every recurring mobility need in the country and be the most trusted name in scheduled rides.",
    },
    "brand_schoolcab": {
        "title": "SchoolCab",
        "body": "SchoolCab is our school commute mobility brand. Its broader aim is to fix the entire school commute ecosystem — reducing the risks associated with school travel and giving parents safety, comfort and complete assurance about their child's daily commute.",
    },
    "brand_cabzee": {
        "title": "Cabzee",
        "body": "Cabzee is our fixed-route office commute brand. It removes the daily stress of booking rides for recurring travel, helps every individual plan their day better, and brings consistency and reliability to the everyday office commute.",
    },
    "brand_advisory": {
        "title": "Mukman Advisory",
        "body": "Mukman Advisory is the group's financial services arm, working on the financial advisory side of mobility platforms — structuring capital, unit economics and growth strategy for the mobility ecosystem.",
    },
    "achievements": {
        "title": "Achievements So Far",
        "body": "In the last four months, Mukman Group has moved from idea to operating business:\n\n- Deployed 17 vehicles to the market\n- Tied up with 3 schools, generating ₹4 lakh per month in recurring revenue from school commute alone\n- Partnered with Milo, a mobility platform, to utilise our vehicle fleet outside school hours\n- On path to generate our first crore of annual revenue\n- Building parallel demand from office-going professionals for their recurring rides\n\nEvery milestone brings us closer to owning the recurring rides category in India.",
        "stats": [
            {"label": "Vehicles Deployed", "value": "17", "sub": "in the last 4 months"},
            {"label": "School Partnerships", "value": "3", "sub": "signed in 4 months"},
            {"label": "Recurring Revenue", "value": "₹4L/mo", "sub": "from school commute"},
            {"label": "Fleet Partner", "value": "Milo", "sub": "off-hours utilisation"},
        ],
    },
    "valuation": {
        "title": "Valuation & Market Opportunity",
        "body": "Mukman Group is currently targeting a valuation of ₹25 crore, reflecting our deployed fleet of 17 vehicles, signed school contracts and a growing recurring revenue base.\n\nJust our home market tells the story: the addressable market for school-student commute in Delhi alone is ₹12,000 crore. Layer on office commuters travelling fixed routes every single day, and the national recurring-rides opportunity runs into tens of thousands of crores — with no organised category leader today.\n\nAs fleet deployment, school tie-ups and office routes scale, we expect this valuation to compound quarter on quarter.",
        "stats": [
            {"label": "Target Valuation", "value": "₹25 Cr", "sub": "current target"},
            {"label": "Delhi School Market", "value": "₹12,000 Cr", "sub": "addressable, Delhi only"},
            {"label": "Vehicles Deployed", "value": "17", "sub": "and scaling"},
            {"label": "Recurring Revenue", "value": "₹4L/mo", "sub": "from school commute"},
        ],
    },
    "benefits": {
        "title": "Employee Benefits at MukMan",
        "body": "At MukMan, we take care of the people who take care of mobility.\n\nMedical Insurance: Every employee is covered with ₹2 lakh of medical insurance from day one — hospitalisation, day-care procedures and emergency treatment — so nobody worries about medical bills while building the future of mobility.\n\nWeekly Refreshment Activities: Every week we step away from the desk together — team activities, games and outings to recharge and bond as one MukMan family.\n\nFounder Counselling Sessions: Every employee gets direct, open-door counselling sessions with the Founder & CEO — career guidance, personal growth, or anything on your mind.\n\nESOP Ownership: Eligible team members hold real ownership in the group through the MukMan ESOP pool. Performance, consistency and discipline directly protect and grow that ownership.\n\nRecognition: Consistent performers are recognised in periodic reviews and all-hands — great work never goes unnoticed at MukMan.",
        "stats": [
            {"label": "Medical Insurance", "value": "₹2 Lakh", "sub": "from day one"},
            {"label": "Refreshment Activities", "value": "Weekly", "sub": "team activities & outings"},
            {"label": "Founder Counselling", "value": "Direct", "sub": "open-door sessions"},
            {"label": "ESOP Pool", "value": "Ownership", "sub": "for eligible members"},
        ],
    },
}

SEED_PENALTIES = [
    {"key": "ts_delay_5d", "label": "Timesheet delayed 5 days", "description": "Continuous delay of 5 days in filling timesheets results in a salary payout decline.", "percent": 10, "applies_to": "salary", "enabled": True},
    {"key": "ts_delay_10d", "label": "Timesheet delayed 10 days", "description": "Continuous failure to fill timesheets for 10 days results in a decline in ownership / ESOP percentage (for ESOP-eligible employees).", "percent": 10, "applies_to": "esop", "enabled": True},
    {"key": "weekly_review_3_missed", "label": "3 weekly summaries missed", "description": "Missing 3 consecutive weekly summaries triggers a penalty on salary or ownership.", "percent": 10, "applies_to": "salary", "enabled": True},
    {"key": "monthly_review_2_missed", "label": "2 monthly summaries missed", "description": "Missing 2 monthly achievement summaries triggers a penalty on salary or ownership.", "percent": 10, "applies_to": "salary", "enabled": True},
]

SEED_EMPLOYEES = [
    {"name": "Abhimanyu", "email": "abhimanyu@mukman.in", "phone": "7807903265", "title": "Operations & Technology", "department": "Operations", "manager": "admin", "esop": True},
    {"name": "Himanshu", "email": "himanshu@mukman.in", "phone": "9560320353", "title": "Fleet Maintenance", "department": "Fleet Maintenance", "manager": "abhimanyu@mukman.in", "esop": False},
    {"name": "Dipanshu", "email": "dipanshu@mukman.in", "phone": "9220836613", "title": "Driver Management", "department": "Driver Operations", "manager": "abhimanyu@mukman.in", "esop": False},
    {"name": "Miki", "email": "micky@mukman.in", "phone": "9311763047", "title": "Collections", "department": "Collections", "manager": "abhimanyu@mukman.in", "esop": False},
    {"name": "Charvi", "email": "charvi@mukman.in", "phone": "9811577867", "title": "Marketing & Parents", "department": "Marketing", "manager": "admin", "esop": False},
    {"name": "Sanjay", "email": "sanjay@mukman.in", "phone": "9060438347", "title": "Technology", "department": "Tech", "manager": "admin", "esop": True},
    {"name": "Ashish", "email": "ashish@mukman.in", "phone": "9354439610", "title": "Attendance, Charging & HSV Global", "department": "Operations Support", "manager": "himanshu@mukman.in", "esop": False, "also": ["dipanshu@mukman.in", "abhimanyu@mukman.in"]},
]

ORG_VIRTUAL = [
    {"key": "driver_hiring_open", "name": "Driver Hiring", "title": "Position open", "kind": "open", "parent_email": "abhimanyu@mukman.in"},
    {"key": "pankaj_finance", "name": "Pankaj", "title": "Finance & Accounts", "kind": "hat", "parent_email": "founder"},
]


@app.on_event("startup")
async def seed():
    await db.users.create_index("email", unique=True)
    await db.cron_runs.create_index("run_id", unique=True)

    admin_email = os.environ.get("ADMIN_EMAIL", "mukman.technologies@gmail.com").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "Mukman@2026")
    admin = await db.users.find_one({"email": admin_email})
    if not admin:
        doc = {"name": "Pankaj Talan", "email": admin_email, "phone": "9000000001",
               "password_hash": hash_password(admin_password), "title": "Founder & CEO",
               "department": "Leadership", "reports_to": None, "role": "admin",
               "esop_applicable": True, "active": True, "created_at": now_iso()}
        res = await db.users.insert_one(doc)
        admin = await db.users.find_one({"_id": res.inserted_id})
    elif not verify_password(admin_password, admin["password_hash"]):
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_password)}})
    await db.users.update_one({"email": admin_email}, {"$set": {"name": "Pankaj Talan", "title": "Founder & CEO"}})
    admin = await db.users.find_one({"email": admin_email})

    for e in SEED_EMPLOYEES:
        if not await db.users.find_one({"email": e["email"]}):
            if e["manager"] == "admin":
                mgr_id = admin["_id"]
            else:
                mgr = await db.users.find_one({"email": e["manager"]})
                mgr_id = mgr["_id"] if mgr else admin["_id"]
            await db.users.insert_one({
                "name": e["name"], "email": e["email"], "phone": e["phone"],
                "password_hash": hash_password("Schoolcab@1"), "title": e["title"],
                "department": e["department"], "reports_to": mgr_id, "role": "employee",
                "esop_applicable": e["esop"], "active": True, "must_change_password": True, "created_at": now_iso(),
            })

    for e in SEED_EMPLOYEES:
        upd = {"name": e["name"], "title": e["title"], "department": e["department"], "phone": e["phone"]}
        if e.get("also"):
            ids = []
            for em in e["also"]:
                m = await db.users.find_one({"email": em})
                if m:
                    ids.append(m["_id"])
            upd["also_reports_to"] = ids
        await db.users.update_one({"email": e["email"]}, {"$set": upd})

    for key, val in SEED_CONTENT.items():
        if not await db.content.find_one({"key": key}):
            await db.content.insert_one({"key": key, **val, "updated_at": now_iso(), "updated_by": admin["_id"]})

    for p in SEED_PENALTIES:
        if not await db.penalty_rules.find_one({"key": p["key"]}):
            await db.penalty_rules.insert_one(p)

    for v in ORG_VIRTUAL:
        if not await db.org_nodes.find_one({"key": v["key"]}):
            await db.org_nodes.insert_one(v)

    logger.info("Seed complete")


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[os.environ.get("FRONTEND_URL", "http://localhost:3000"), "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
